import streamlit as st
import pandas as pd
from PyPDF2 import PdfReader
import json
import xml.etree.ElementTree as ET
from docx import Document
import os
from fpdf import FPDF
import io
from datetime import datetime
from textblob import TextBlob


# Função para ler arquivos XLSX e transformar em TXT
def read_xlsx(file):
    # Define o caminho para o arquivo TXT
    caminho_arquivo = f'./conhecimento/{os.path.splitext(os.path.basename(file.name))[0]}.txt'

    # Abre o arquivo XLSX e lê todas as abas
    with pd.ExcelFile(file) as xls:
        with open(caminho_arquivo, 'w', encoding='utf-8') as txt_file:
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name)
                # Escreve o nome da aba no arquivo TXT
                txt_file.write(f'--- Aba: {sheet_name} ---\n')
                # Salva o DataFrame como texto, usando tabulação como delimitador
                df.to_csv(txt_file, sep='\t', index=False, header=True)  # Inclui cabeçalho
                txt_file.write('\n')  # Adiciona uma nova linha após cada aba

    # Lê o arquivo TXT e retorna seu conteúdo
    with open(caminho_arquivo, 'r', encoding='utf-8') as arquivo:
        info = arquivo.read()
    return info


# Função para ler arquivos PDF e transformar em TXT
def read_pdf(file):
    text = ""
    # Lê o arquivo PDF
    pdf_reader = PdfReader(file)
    for page in pdf_reader.pages:
        text += page.extract_text()
    # Define o caminho para o arquivo TXT
    caminho_arquivo = f'./conhecimento/{os.path.splitext(os.path.basename(file.name))[0]}.txt'
    # Salva o texto extraído como um arquivo TXT
    with open(caminho_arquivo, 'w', encoding='utf-8') as txt_file:
        txt_file.write(text)
    # Retorna o texto extraído
    return text


# Função para ler arquivos JSON e transformar em TXT
def read_json(file):
    # Lê o arquivo JSON
    content = json.load(file)
    # Define o caminho para o arquivo TXT
    caminho_arquivo = f'./conhecimento/{os.path.splitext(os.path.basename(file.name))[0]}.txt'
    # Salva o conteúdo JSON como um arquivo TXT
    with open(caminho_arquivo, 'w', encoding='utf-8') as txt_file:
        txt_file.write(json.dumps(content, indent=4))  # Formata o conteúdo JSON com indentação
    # Retorna o conteúdo lido do JSON
    return content


# Função para ler arquivos XML e transformar em TXT
def read_xml(file):
    # Lê o arquivo XML
    tree = ET.parse(file)
    root = tree.getroot()
    # Converte o conteúdo XML em uma string
    xml_str = ET.tostring(root, encoding='utf-8').decode('utf-8')
    # Define o caminho para o arquivo TXT
    caminho_arquivo = f'./conhecimento/{os.path.splitext(os.path.basename(file.name))[0]}.txt'
    # Salva o conteúdo XML como um arquivo TXT
    with open(caminho_arquivo, 'w', encoding='utf-8') as txt_file:
        txt_file.write(xml_str)
    # Retorna a string XML
    return xml_str


# Função para ler arquivos HTML e transformar em TXT
def read_html(file):
    # Lê o conteúdo do arquivo HTML e decodifica como UTF-8
    html_content = file.read().decode("utf-8")
    # Define o caminho para o arquivo TXT
    caminho_arquivo = f'./conhecimento/{os.path.splitext(os.path.basename(file.name))[0]}.txt'
    # Salva o conteúdo HTML como um arquivo TXT
    with open(caminho_arquivo, 'w', encoding='utf-8') as txt_file:
        txt_file.write(html_content)
    # Retorna o conteúdo HTML
    return html_content


# Função para ler arquivos DOCX e transformar em TXT
def read_docx(file):
    # Lê o arquivo DOCX
    doc = Document(file)
    text = ""
    # Extrai texto de cada parágrafo
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"  # Adiciona uma nova linha após cada parágrafo
    # Define o caminho para o arquivo TXT
    caminho_arquivo = f'./conhecimento/{os.path.splitext(os.path.basename(file.name))[0]}.txt'
    # Salva o texto extraído como um arquivo TXT
    with open(caminho_arquivo, 'w', encoding='utf-8') as txt_file:
        txt_file.write(text)
    # Retorna o texto extraído
    return text


# Função para ler arquivos TXT
def read_txt(file):
    return file.read().decode("utf-8")


def carregar_arquivos():
    # Sidebar para carregamento de arquivos
    st.sidebar.header("Carregar Documentos")
    # Apresentando o botão para carregar os documentos
    with st.sidebar:
        st.subheader('Clique no botão abaixo para carregar seus dados e fazer uma análise com o Doctor Med:')
        # Carregar múltiplos arquivos
        uploaded_files = st.sidebar.file_uploader("Coloque seu arquivo aqui:", type=["xlsx", "pdf", "xml", "json", "html", "htm", "doc", "docx", "txt", "xls"], accept_multiple_files=True)

        if st.button('CARREGAR'):
            conteudos = []  # Lista para armazenar o conteúdo dos arquivos
            for file in uploaded_files:
                with st.spinner("Processing"):
                    st.write(f"**Arquivo carregado :** {file.name}")

                    if file.type in ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel"]:
                        conteudo = read_xlsx(file)
                        conteudos.append(conteudo)

                    elif file.type == "application/pdf":
                        conteudo = read_pdf(file)
                        conteudos.append(conteudo)

                    elif file.type == "application/json":
                        conteudo = read_json(file)
                        conteudos.append(json.dumps(conteudo, indent=4))  # Formata o JSON

                    elif file.type in ["application/xml", "text/xml"]:
                        conteudo = read_xml(file)
                        conteudos.append(conteudo)

                    elif file.type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"]:
                        conteudo = read_docx(file)
                        conteudos.append(conteudo)

                    elif file.type == "text/plain":
                        # Lendo arquivos de texto (.txt)
                        caminho_arquivo = file.name  # Usando o nome do arquivo para leitura
                        with open(caminho_arquivo, 'r', encoding='utf-8') as arquivo:
                            conteudo = arquivo.read()  # Lê o conteúdo do arquivo
                            conteudos.append(conteudo)

                    elif file.type in ["text/html", "text/htm"]:
                        conteudo = read_html(file)
                        conteudos.append(conteudo)

            return conteudos  # Retorna a lista de conteúdos


# Função para salvar a imagem no servidor
def save_uploaded_file(uploaded_file, user_type):
    """Salva um arquivo de imagem no diretório correspondente ao tipo de usuário.

    Args:
        uploaded_file: O arquivo de imagem carregado.
        user_type (str): O tipo de usuário (admin, parceiro, colaborador ou cliente).

    Returns:
        str: O caminho do arquivo salvo.
    """

    # Define o diretório base onde as imagens serão salvas
    base_dir = "MEDIA"

    # Mapeia o tipo de usuário para o diretório correspondente
    user_dirs = {
        "Admin": "admin",
        "Parceiro": "parceiro",
        "Colaborador": "colaborador",
        "Cliente": "cliente"
    }

    # Verifica se o tipo de usuário é válido
    if user_type not in user_dirs:
        raise ValueError(f"Tipo de usuário inválido: {user_type}")

    # Cria o caminho final onde a imagem será salva
    user_dir = os.path.join(base_dir, user_dirs[user_type])

    # Cria o diretório se não existir
    if not os.path.exists(user_dir):
        os.makedirs(user_dir)

    # Salva o arquivo na pasta definida
    file_path = os.path.join(user_dir, uploaded_file.name)
    with open(file_path, "wb") as f:
        f.write(uploaded_file.getbuffer())

    return file_path


# Função para gerar o PDF
def create_pdf(messages):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    role = messages["role"].capitalize()
    content = messages["content"]
    pdf.cell(200, 10, txt=f"{role}: {content}", ln=True)
    pdf_bytes = create_pdf(st.session_state.messages)
    st.download_button(
        label="Baixar PDF",
        data=pdf_bytes,
        file_name="conversa.pdf",
        mime="application/pdf",
    )

    return pdf.output(dest='S').encode('latin1')


# Função para gerar o Excel
def create_excel(messages):
    df = pd.DataFrame(messages)
    buffer = io.BytesIO()
    df.to_excel(buffer, index=False)
    buffer.seek(0)
    st.download_button(
        label="Baixar Excel",
        data=df,
        file_name="conversa.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    st.stop()

    return buffer.getvalue()


# Função global para formatar data e hora
def formatar_data_hora(valor):
    if isinstance(valor, str):  # Verifica se é uma string
        try:
            # Tenta converter a string para datetime
            date_obj = datetime.strptime(valor, '%Y-%m-%d %H:%M:%S')  # Formato esperado
            return date_obj.strftime('%d/%m/%Y'), date_obj.strftime('%H:%M:%S')
        except ValueError:
            pass  # Ignora erro de formato inválido
    elif isinstance(valor, datetime):  # Verifica se é um objeto datetime
        # Formata a data e a hora
        return valor.strftime('%d/%m/%Y'), valor.strftime('%H:%M:%S')
    return None  # Retorna None se o valor for inválido ou vazio


vantagens_oraculo_biblia = """
        O Oráculo Bíblia oferece uma série de funcionalidades que podem transformar a gestão de uma igreja, tornando-a mais 
        eficiente e organizada. Abaixo estão as explicações detalhadas sobre cada uma dessas funcionalidades, destacando suas 
        vantagens e benefícios.

        1. Gestão de Membros:
           A funcionalidade de gestão de membros permite que as igrejas mantenham um registro atualizado de todos os seus 
           integrantes. Isso significa que você pode facilmente acessar informações como nome, contato, histórico de 
           participação e até mesmo dados sobre a contribuição de cada membro. Com isso, é possível promover um acompanhamento 
           mais próximo e personalizado, fortalecendo a comunidade e facilitando o envolvimento dos membros nas atividades da 
           igreja.

        2. Gerenciamento de Departamentos:
           Com o gerenciamento de departamentos, as igrejas podem organizar suas atividades em diferentes áreas, como 
           ministérios, grupos de jovens, e serviços sociais. Isso ajuda a distribuir responsabilidades de forma clara e a 
           garantir que cada departamento tenha os recursos e o suporte necessários para funcionar efetivamente. Além disso, 
           facilita a comunicação entre os departamentos, promovendo uma colaboração mais eficaz .

        3. Gestão de Cultos:
           A gestão de cultos permite que as igrejas planejem e organizem seus serviços de maneira mais eficiente. É possível 
           programar datas, horários, e temas dos cultos, além de gerenciar a participação dos membros e a logística necessária.
            Isso garante que tudo esteja preparado para oferecer a melhor experiência aos frequentadores, tornando os cultos 
            mais impactantes e bem organizados.

        4. Gerenciamento de Enquetes:
           A funcionalidade de gerenciamento de enquetes possibilita que as igrejas coletem feedback dos membros de maneira 
           simples e rápida. Isso pode ser útil para entender as preferências da congregação sobre temas, horários de cultos, e 
           outras atividades. Ao ouvir a voz dos membros, a igreja pode ajustar suas programações e iniciativas para atender 
           melhor às necessidades da comunidade.

        5. Gerenciamento de Eventos:
           Com o gerenciamento de eventos, as igrejas podem planejar atividades especiais, como retiros, conferências, e ações 
           comunitárias. Essa funcionalidade ajuda a manter um calendário organizado, controla inscrições e promove uma 
           comunicação clara sobre os eventos, garantindo que os membros estejam sempre informados e engajados.

        6. Gestão de Dízimos e Ofertas:
           A gestão de dízimos e ofertas oferece uma maneira prática de monitorar as contribuições financeiras dos membros. Isso
            não apenas facilita a contabilidade da igreja, mas também ajuda a identificar padrões de doação, permitindo que a 
            liderança da igreja planeje melhor suas finanças e projetos. Além disso, isso pode promover uma maior transparência 
            e confiança entre os membros.

        7. API com Sistema de Pagamentos (ASAAS):
           Com a integração de uma API que permite transações financeiras seguras, as igrejas podem simplificar o processo de 
           recebimento de dízimos e ofertas. Isso proporciona uma experiência prática para os membros, que podem contribuir de 
           forma rápida e segura, seja presencialmente ou online. Esse tipo de facilidade pode aumentar a participação dos 
           membros nas contribuições.

        8. Análise Inteligente com Modelo LLM:
           A análise inteligente utiliza modelos de linguagem para processar e interpretar dados de maneira mais eficaz. Isso 
           permite que as igrejas obtenham insights valiosos sobre a participação dos membros, padrões de doação, e outras 
           métricas relevantes. Com esses dados, a liderança pode tomar decisões mais informadas sobre futuras iniciativas e 
           programas.

        9. Automação com Make Automation - Disparo de Emails e Agendamentos:
            A automação é um recurso poderoso que pode economizar tempo e esforço. Com a automação, as igrejas podem programar 
            envios de e-mails, lembretes de eventos, e agendamentos, tudo de forma automática. Isso garante que os membros 
            estejam sempre informados e ajuda a manter a comunicação fluida, sem a necessidade de esforços manuais.
            
            Por que as igrejas devem ter um sistema de gestão eclesiástica como este?
            Em resumo, um sistema de gestão eclesiástica como o oferecido pelo Oráculo Bíblia permite que as igrejas operem de forma
            mais eficiente, economizando tempo e recursos. Ele promove uma melhor organização, comunicação e engajamento entre os 
            membros, facilitando a gestão de atividades e a transparência financeira. Ao adotar esse sistema, as igrejas podem 
            focar em sua missão principal de servir à comunidade, enquanto desfrutam de uma gestão interna mais estruturada e 
            eficaz.
        
        10. Gerenciamento e Gestão de Lojas e Cadastro de Produtos:
            Esta funcionalidade é ideal para igrejas que possuem lojas internas para comercializar produtos cristãos. 
            Com o gerenciamento de lojas, as igrejas podem cadastrar produtos, controlar estoque e registrar vendas de 
            maneira organizada. Isso não só facilita a operação da loja, mas também promove um ambiente mais eficiente 
            para a venda de itens que são úteis para a comunidade, como livros, materiais de estudo e produtos 
            religiosos. Além disso, um sistema de gestão permite que a igreja analise as vendas e o desempenho dos 
            produtos, ajudando na tomada de decisões sobre quais itens oferecer e como melhor atender os anseios dos 
            membros da congregação
        
            """


historia_recomecar = """

            Em 2005, o pr. Marciano e sua esposa, Maria dos Santos, pais do pr. Ronaldo Santos, inauguram a Igreja do 
            Deus Todo-Poderoso na Avenida José Faria da Rocha, no bairro Eldorado, em Contagem/MG. Em 2006, o 
            pr. Ronaldo Santos é consagrado pastor na mesma igreja. No mesmo ano, sob direção divina, ele assume a 
            direção da igreja e altera o nome para Comunidade Cristã Recomeçar, transferindo sua sede para um imóvel 
            maior, ainda na Avenida José Faria da Rocha, no bairro Eldorado. 

            Em 2017, o pr. Ronaldo Santos, com a ajuda de sua esposa, pra. Sanmya Santos, decide dar um passo de fé e 
            transfere a sede da igreja para o atual endereço, a antiga sede da fábrica de roupas Radial, na 
            Av. Dr. João Augusto Fonseca e Silva, 387, no bairro Novo Eldorado, em Contagem/MG. Hoje, a Comunidade 
            Cristã Recomeçar é uma grande família em constante crescimento, buscando sempre o arrependimento, 
            relacionamento de intimidade com Deus através da oração, leitura das Escrituras e serviço ao próximo, 
            visando à plenitude da vida cristã. 
            
            
            A Comunidade Cristã Recomeçar fica na Av. Dr. João Augusto Fonseca e Silva, 387, no bairro Novo Eldorado, 
            em Contagem/MG. Toque aqui para acessar no mapa.
            
            """

localizacao = 'https://maps.app.goo.gl/TTvz9U8PSMEinH6t5'

missao_visao_valores = """

            Missão: Nosso propósito é manifestar o Reino de Deus às pessoas, amando e servindo ao próximo, com o 
            objetivo de tornar Jesus conhecido em todas as nações e conduzir as pessoas a um relacionamento íntimo 
            com Deus.
            
            Visão: Ser uma fonte de restauração para as pessoas em todas as esferas da cidade de Contagem.
            
            Valores: Unidade, Fé, Atitude e Amor ao próximo são os pilares que norteiam nosso trabalho e orientam 
            nossa conduta diária.
            
"""


class DizimoOferta:
    def __init__(self):
        # Links para dízimos e ofertas
        self.links = {
            "dz_domingo_manha": "https://sandbox.asaas.com/c/qt9n57f0y2473pet",
            "of_domingo_manha": "https://sandbox.asaas.com/c/5h9pqmn4rxsn2cml",
            "dz_domingo_noite": "https://sandbox.asaas.com/c/xo23koleskzuwosr",
            "of_domingo_noite": "https://sandbox.asaas.com/c/xo23koleskzuwosr",
            "dz_quinta": "https://sandbox.asaas.com/c/bw5rcpapxmuu4umo",
            "of_quinta": "https://sandbox.asaas.com/c/i4p70kwrdudodblc"
        }

    def obter_link(self, tipo, culto):
        """
        Retorna o link correspondente ao tipo (dízimo/oferta) e culto.
        """
        if tipo == "dizimo":
            if culto == "domingo_manha":
                return self.links["dz_domingo_manha"]
            elif culto == "domingo_noite":
                return self.links["dz_domingo_noite"]
            elif culto == "quinta_noite":
                return self.links["dz_quinta"]
        elif tipo == "oferta":
            if culto == "domingo_manha":
                return self.links["of_domingo_manha"]
            elif culto == "domingo_noite":
                return self.links["of_domingo_noite"]
            elif culto == "quinta_noite":
                return self.links["of_quinta"]
        return None

    def processar_intencao(self, mensagem):
        """
        Processa a intenção do usuário com base nas palavras-chave.
        """
        mensagem_lower = mensagem.lower()

        # Verifica se o usuário quer um link
        for palavra in PALAVRAS_PERMITIDAS["link"]:
            if palavra in mensagem_lower:
                return "Por favor, especifique se você quer o link de dízimo ou oferta e o culto em que está participando."

        # Verifica se o usuário quer dar uma oferta
        for palavra in PALAVRAS_PERMITIDAS["oferta"]:
            if palavra in mensagem_lower:
                return "Você deseja dar uma oferta. Por favor, informe o culto em que está participando."

        # Verifica se o usuário quer devolver o dízimo
        for palavra in PALAVRAS_PERMITIDAS["dizimo"]:
            if palavra in mensagem_lower:
                return "Você deseja devolver seu dízimo. Por favor, informe o culto em que está participando."

        # Identifica o culto mencionado pelo usuário
        for culto, palavras in PALAVRAS_PERMITIDAS["culto"].items():
            for palavra in palavras:
                if palavra in mensagem_lower:
                    return f"Você está participando do culto de {culto.replace('_', ' ')}. Posso ajudar com o link de dízimo ou oferta?"

        # Caso nenhuma intenção seja identificada
        return "Desculpe, não entendi sua solicitação. Você pode pedir links de dízimo/oferta ou informar seu culto."